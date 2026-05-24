import csv
import io
import json
import math
import os
import re
import tempfile
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
import cgi
from urllib.parse import quote

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
UPLOAD_DIR = BASE_DIR / "uploads"
EXPORT_DIR = BASE_DIR / "exports"
UPLOAD_DIR.mkdir(exist_ok=True)
EXPORT_DIR.mkdir(exist_ok=True)

MAX_PREVIEW_ROWS = 20
MAX_RESULT_ROWS = 500
HEADER_SCAN_ROWS = 12


DEFAULT_FIELDS = {
    "sap": {
        "material": ["物料", "SKU编码", "SKU", "商品编码"],
        "quantity": ["数量", "变更数量", "移动数量"],
        "posting_date": ["过账日期", "输入日期", "变动日期"],
        "external_doc": ["外部系统单号", "原始单据号(记录自动产生的单据的源头单据)", "参照"],
        "unloading_point": ["卸货点"],
    },
    "logic": {
        "material": ["SKU编码", "物料", "SKU", "商品编码"],
        "quantity": ["变更数量", "数量", "移动数量"],
        "change_datetime": ["变动时间"],
        "doc_type": ["单据类型"],
        "pre_doc": ["前置单据编码"],
        "receipt_doc": ["收发货单据"],
        "warehouse": ["逻辑仓编码", "库存地点", "仓库", "仓库编码"],
    },
}

GROUP1_LOGIC_TYPES = [
    "零售订单",
    "零售退单",
    "JIT唯品退供单",
    "JIT唯品配货单",
    "零售发货通知单",
    "零售退货通知单",
    "平台寄售",
    "平台退供",
    "shein",
    "拼多多TEMU",
]
GROUP1_EXPORT_DOC_KEYS = ["BMS", "ARST", "RO", "RC"]
GROUP2_LOGIC_TYPES = ["其他出入库单", "库存调整单"]
GROUP3_LOGIC_TYPES = ["样衣调出", "样衣领用", "样衣归还"]
GROUP4_LOGIC_TYPES = ["销售出库单", "销售退货单", "销售退单", "返修出库", "采购交货单"]
GROUP4_EXPORT_DOC_KEYS = ["XSD", "XSTD"]
GROUP5_LOGIC_TYPES = ["采购退货单"]


def normalize_header(value):
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip()).lower()


def normalize_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_number(value):
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if math.isnan(value):
            return Decimal("0")
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    if not text:
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation:
        match = re.search(r"-?\d+(?:\.\d+)?", text)
        return Decimal(match.group(0)) if match else Decimal("0")


def decimal_to_json(value):
    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return int(value)
        return float(value)
    return value


def serialize_cell(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, Decimal):
        return decimal_to_json(value)
    if value is None:
        return ""
    return value


def parse_datetime_value(date_value, time_value=None):
    if isinstance(date_value, datetime):
        base = date_value
    elif isinstance(date_value, date):
        base = datetime.combine(date_value, time.min)
    else:
        text = normalize_text(date_value)
        if not text:
            return ""
        parsed = None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            return text
        base = parsed

    if time_value:
        if isinstance(time_value, time):
            base = datetime.combine(base.date(), time_value)
        else:
            text = normalize_text(time_value)
            for fmt in ("%H:%M:%S", "%H:%M"):
                try:
                    parsed_time = datetime.strptime(text, fmt).time()
                    base = datetime.combine(base.date(), parsed_time)
                    break
                except ValueError:
                    continue
    return base.strftime("%Y-%m-%d %H:%M:%S")


def safe_filename(name):
    cleaned = re.sub(r"[^\w\u4e00-\u9fff.\- ]+", "_", name).strip()
    return cleaned or "upload.xlsx"


def detect_header(rows):
    best_index = 0
    best_score = -1
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        values = [normalize_text(v) for v in row]
        non_empty = sum(1 for v in values if v)
        unique = len(set(v for v in values if v))
        score = non_empty + unique
        if any(v in ("物料", "SKU编码", "变更数量", "数量") for v in values):
            score += 20
        if score > best_score:
            best_index = idx
            best_score = score
    return best_index


def read_workbook(path, preview_only=False):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    scanned = []
    for _ in range(HEADER_SCAN_ROWS):
        try:
            scanned.append(next(rows_iter))
        except StopIteration:
            break
    header_index = detect_header(scanned)
    header_row = scanned[header_index] if scanned else []
    headers = []
    counts = Counter()
    for idx, value in enumerate(header_row):
        label = normalize_text(value) or f"空列{idx + 1}"
        counts[label] += 1
        if counts[label] > 1:
            label = f"{label}_{counts[label]}"
        headers.append(label)

    rows = []
    source_rows = scanned[header_index + 1 :]
    if not preview_only:
        source_rows = list(source_rows) + list(rows_iter)

    limit = MAX_PREVIEW_ROWS if preview_only else None
    for row_number, row in enumerate(source_rows, start=header_index + 2):
        if limit and len(rows) >= limit:
            break
        values = list(row[: len(headers)])
        if not any(normalize_text(v) for v in values):
            continue
        record = {headers[i]: serialize_cell(values[i]) if i < len(values) else "" for i in range(len(headers))}
        record["_row_number"] = row_number
        rows.append(record)

    return {
        "sheet": sheet.title,
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "header_row": header_index + 1,
        "headers": headers,
        "preview": rows[:MAX_PREVIEW_ROWS],
        "rows": rows if not preview_only else [],
    }


def guess_mapping(headers, side):
    normalized = {normalize_header(header): header for header in headers}
    mapping = {}
    for field, candidates in DEFAULT_FIELDS[side].items():
        selected = ""
        for candidate in candidates:
            key = normalize_header(candidate)
            if key in normalized:
                selected = normalized[key]
                break
        if not selected:
            for header in headers:
                header_key = normalize_header(header)
                if any(normalize_header(candidate) and normalize_header(candidate) in header_key for candidate in candidates):
                    selected = header
                    break
        mapping[field] = selected
    return mapping


def upload_file(file_item):
    filename = safe_filename(file_item.filename or "upload.xlsx")
    suffix = Path(filename).suffix or ".xlsx"
    upload_id = uuid.uuid4().hex
    path = UPLOAD_DIR / f"{upload_id}{suffix}"
    with path.open("wb") as fh:
        fh.write(file_item.file.read())
    workbook = read_workbook(path, preview_only=True)
    side = "logic" if any(h in workbook["headers"] for h in ("SKU编码", "变更数量", "逻辑仓编码")) else "sap"
    return {
        "id": upload_id,
        "name": filename,
        "path": str(path),
        "side_guess": side,
        "workbook": workbook,
        "mapping_guess": guess_mapping(workbook["headers"], side),
    }


def contains_any(value, keywords):
    text = normalize_text(value).lower()
    return any(str(keyword).lower() in text for keyword in keywords)


def not_contains_any(value, keywords):
    return not contains_any(value, keywords)


def field_value(row, fields, field):
    return normalize_text(row.get(fields.get(field, ""), ""))


def field_quantity(row, fields):
    return normalize_number(row.get(fields.get("quantity", ""), ""))


def field_day(row, fields, field):
    value = row.get(fields.get(field, ""), "")
    parsed = parse_datetime_value(value)
    if re.match(r"\d{4}-\d{2}-\d{2}", parsed):
        return parsed[:10]
    return normalize_text(value)[:10]


def logic_day(row, fields):
    return field_day(row, fields, "change_datetime")


def sap_day(row, fields):
    return field_day(row, fields, "posting_date")


def rr_or_sr_text(value):
    return bool(re.search(r"\b(?:RR|SR)[A-Za-z0-9_-]*", normalize_text(value), re.IGNORECASE))


def extract_rr_doc(value):
    match = re.search(r"\bRR[A-Za-z0-9_-]*", normalize_text(value), re.IGNORECASE)
    return match.group(0).upper() if match else ""


def normalize_doc(value):
    return normalize_text(value).upper()


def group_sum(rows, fields, date_func, doc_field=None):
    grouped = defaultdict(lambda: {"quantity": Decimal("0"), "rows": 0, "docs": set()})
    for row in rows:
        material = field_value(row, fields, "material")
        day = date_func(row, fields)
        if not material or not day:
            continue
        key = (material, day)
        grouped[key]["quantity"] += field_quantity(row, fields)
        grouped[key]["rows"] += 1
        if doc_field:
            doc = field_value(row, fields, doc_field)
            if doc:
                grouped[key]["docs"].add(doc)
    return grouped


def compare_date_group(
    group_name,
    logic_rows,
    sap_rows,
    logic_fields,
    sap_fields,
    logic_doc_field=None,
    sap_doc_field=None,
    include_counts=True,
):
    logic_grouped = group_sum(logic_rows, logic_fields, logic_day, logic_doc_field)
    sap_grouped = group_sum(sap_rows, sap_fields, sap_day, sap_doc_field)
    result = []
    all_keys = sorted(set(logic_grouped) | set(sap_grouped))
    for material, day in all_keys:
        logic_bucket = logic_grouped.get((material, day), {})
        sap_bucket = sap_grouped.get((material, day), {})
        logic_qty = logic_bucket.get("quantity", Decimal("0"))
        sap_qty = sap_bucket.get("quantity", Decimal("0"))
        diff = logic_qty - sap_qty
        if diff == 0:
            continue
        row = {
            "分组": group_name,
            "SKU/物料": material,
            "日期": day,
            "逻辑仓数量": decimal_to_json(logic_qty),
            "EXPORT数量": decimal_to_json(sap_qty),
            "差异数量": decimal_to_json(diff),
        }
        if logic_doc_field:
            row["逻辑仓前置单据编码"] = "、".join(sorted(logic_bucket.get("docs", set())))
        if sap_doc_field:
            row["SAP外部系统单号"] = "、".join(sorted(sap_bucket.get("docs", set())))
        if include_counts:
            row["逻辑仓行数"] = logic_bucket.get("rows", 0)
            row["EXPORT行数"] = sap_bucket.get("rows", 0)
        result.append(row)
    result.sort(key=lambda row: (row["SKU/物料"], row["日期"]))
    return result


def missing_docs_from_export(
    group_name,
    logic_rows,
    sap_rows,
    logic_fields,
    sap_fields,
    logic_doc_field,
    sap_doc_field="external_doc",
    doc_transform=None,
    include_original_doc=True,
    include_logic_row=True,
):
    sap_docs = {
        (
            field_value(row, sap_fields, "material"),
            doc_transform(field_value(row, sap_fields, sap_doc_field))
            if doc_transform
            else normalize_doc(field_value(row, sap_fields, sap_doc_field)),
        )
        for row in sap_rows
        if (
            doc_transform(field_value(row, sap_fields, sap_doc_field))
            if doc_transform
            else normalize_doc(field_value(row, sap_fields, sap_doc_field))
        )
    }
    missing = []
    seen = set()
    for row in logic_rows:
        doc = field_value(row, logic_fields, logic_doc_field)
        doc_key = doc_transform(doc) if doc_transform else normalize_doc(doc)
        if not doc_key:
            continue
        material = field_value(row, logic_fields, "material")
        key = (material, doc_key)
        if key in sap_docs or key in seen:
            continue
        seen.add(key)
        result_row = {
            "分组": group_name,
            "SKU/物料": material,
            "逻辑仓单号": doc_key,
            "单据类型": field_value(row, logic_fields, "doc_type"),
            "逻辑仓编码": field_value(row, logic_fields, "warehouse"),
            "变动日期": logic_day(row, logic_fields),
            "变更数量": decimal_to_json(field_quantity(row, logic_fields)),
            "结果": "EXPORT不存在",
        }
        if include_original_doc:
            result_row["原始单号"] = doc
        if include_logic_row:
            result_row["逻辑仓行号"] = row.get("_row_number", "")
        missing.append(result_row)
    missing.sort(key=lambda row: (row["SKU/物料"], row["逻辑仓单号"]))
    return missing


def compare_records(sap_rows, logic_rows, sap_fields, logic_fields):
    logic_g1 = [
        row
        for row in logic_rows
        if contains_any(field_value(row, logic_fields, "doc_type"), GROUP1_LOGIC_TYPES)
    ]
    sap_g1 = [
        row
        for row in sap_rows
        if contains_any(field_value(row, sap_fields, "external_doc"), GROUP1_EXPORT_DOC_KEYS)
    ]

    logic_g2 = [
        row
        for row in logic_rows
        if (
            contains_any(field_value(row, logic_fields, "doc_type"), ["其他出入库单"])
            and field_value(row, logic_fields, "warehouse") != "L014"
        )
        or contains_any(field_value(row, logic_fields, "doc_type"), ["库存调整单"])
    ]

    logic_g3 = [
        row
        for row in logic_rows
        if contains_any(field_value(row, logic_fields, "doc_type"), GROUP3_LOGIC_TYPES)
        and extract_rr_doc(field_value(row, logic_fields, "receipt_doc"))
        and field_value(row, logic_fields, "warehouse") != "L006"
    ]
    sap_g3 = [
        row
        for row in sap_rows
        if not field_value(row, sap_fields, "external_doc")
        and rr_or_sr_text(field_value(row, sap_fields, "unloading_point"))
    ]

    logic_g4 = [
        row
        for row in logic_rows
        if contains_any(field_value(row, logic_fields, "doc_type"), GROUP4_LOGIC_TYPES)
    ]
    sap_g4 = [
        row
        for row in sap_rows
        if (
            not field_value(row, sap_fields, "external_doc")
            and not rr_or_sr_text(field_value(row, sap_fields, "unloading_point"))
        )
        or (
            field_value(row, sap_fields, "external_doc")
            and contains_any(field_value(row, sap_fields, "external_doc"), GROUP4_EXPORT_DOC_KEYS)
        )
    ]

    logic_g5 = [
        row
        for row in logic_rows
        if contains_any(field_value(row, logic_fields, "doc_type"), GROUP5_LOGIC_TYPES)
    ]

    group1 = compare_date_group(
        "分组一",
        logic_g1,
        sap_g1,
        logic_fields,
        sap_fields,
        logic_doc_field="pre_doc",
        sap_doc_field="external_doc",
        include_counts=False,
    )
    group2 = missing_docs_from_export(
        "分组二",
        logic_g2,
        sap_rows,
        logic_fields,
        sap_fields,
        "pre_doc",
        include_original_doc=False,
        include_logic_row=False,
    )
    group3 = missing_docs_from_export(
        "分组三",
        logic_g3,
        sap_g3,
        logic_fields,
        sap_fields,
        "receipt_doc",
        sap_doc_field="unloading_point",
        doc_transform=extract_rr_doc,
        include_original_doc=False,
        include_logic_row=False,
    )
    group4 = compare_date_group(
        "分组四",
        logic_g4,
        sap_g4,
        logic_fields,
        sap_fields,
        logic_doc_field="pre_doc",
        include_counts=False,
    )
    group5 = missing_docs_from_export(
        "分组五",
        logic_g5,
        sap_rows,
        logic_fields,
        sap_fields,
        "receipt_doc",
        include_original_doc=False,
        include_logic_row=False,
    )

    return {
        "counts": {
            "sap_rows": len(sap_rows),
            "logic_rows": len(logic_rows),
            "group1": len(group1),
            "group2": len(group2),
            "group3": len(group3),
            "group4": len(group4),
            "group5": len(group5),
        },
        "filters": {
            "group1_logic_rows": len(logic_g1),
            "group1_export_rows": len(sap_g1),
            "group2_logic_rows": len(logic_g2),
            "group3_logic_rows": len(logic_g3),
            "group3_export_rows": len(sap_g3),
            "group4_logic_rows": len(logic_g4),
            "group4_export_rows": len(sap_g4),
            "group5_logic_rows": len(logic_g5),
        },
        "group1": group1[:MAX_RESULT_ROWS],
        "group2": group2[:MAX_RESULT_ROWS],
        "group3": group3[:MAX_RESULT_ROWS],
        "group4": group4[:MAX_RESULT_ROWS],
        "group5": group5[:MAX_RESULT_ROWS],
    }


def make_csv(rows):
    output = io.StringIO()
    if not rows:
        return ""
    headers = list(rows[0].keys())
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        safe_row = {}
        for key, value in row.items():
            if isinstance(value, (dict, list, tuple)):
                safe_row[key] = json.dumps(value, ensure_ascii=False)
            else:
                safe_row[key] = value
        writer.writerow(safe_row)
    return output.getvalue()


class Handler(BaseHTTPRequestHandler):
    server_version = "InventoryCompare/0.1"

    def send_json(self, payload, status=HTTPStatus.OK):
        body = json.dumps(payload, ensure_ascii=False, default=decimal_to_json).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_text(self, text, content_type="text/plain; charset=utf-8", status=HTTPStatus.OK, filename=None):
        body = text.encode("utf-8-sig" if content_type.startswith("text/csv") else "utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        if filename:
            encoded = quote(filename)
            self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{encoded}")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format, *args):
        return

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self.serve_static("index.html")
            return
        if parsed.path.startswith("/static/"):
            self.serve_static(parsed.path.removeprefix("/static/"))
            return
        if parsed.path in ("/app.js", "/styles.css"):
            self.serve_static(parsed.path.removeprefix("/"))
            return
        if parsed.path == "/api/export":
            query = parse_qs(parsed.query)
            self.handle_export(query.get("table", ["group1"])[0])
            return
        self.send_error(HTTPStatus.NOT_FOUND)

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/upload":
            self.handle_upload()
            return
        if parsed.path == "/api/compare":
            self.handle_compare()
            return
        if parsed.path == "/api/export":
            payload = self.read_json_body()
            self.handle_export(payload.get("table", "group1"))
            return
        if parsed.path == "/api/save-export":
            payload = self.read_json_body()
            self.handle_save_export(payload.get("table", "group1"))
            return
        self.send_error(HTTPStatus.NOT_FOUND)

    def serve_static(self, relative_path):
        path = (STATIC_DIR / relative_path).resolve()
        if not str(path).startswith(str(STATIC_DIR.resolve())) or not path.exists():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        content_types = {
            ".html": "text/html; charset=utf-8",
            ".css": "text/css; charset=utf-8",
            ".js": "application/javascript; charset=utf-8",
        }
        body = path.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_types.get(path.suffix, "application/octet-stream"))
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_upload(self):
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST"})
        file_item = form["file"] if "file" in form else None
        if file_item is None or not getattr(file_item, "filename", ""):
            self.send_json({"error": "没有收到文件"}, HTTPStatus.BAD_REQUEST)
            return
        try:
            result = upload_file(file_item)
            self.send_json(result)
        except Exception as exc:
            self.send_json({"error": f"读取文件失败：{exc}"}, HTTPStatus.BAD_REQUEST)

    def read_json_body(self):
        length = int(self.headers.get("Content-Length", "0"))
        return json.loads(self.rfile.read(length).decode("utf-8"))

    def load_uploaded_rows(self, upload_id):
        matches = list(UPLOAD_DIR.glob(f"{upload_id}.*"))
        if not matches:
            raise FileNotFoundError("找不到上传文件")
        return read_workbook(matches[0], preview_only=False)["rows"]

    def handle_compare(self):
        try:
            payload = self.read_json_body()
            sap_rows = self.load_uploaded_rows(payload["sap"]["id"])
            logic_rows = self.load_uploaded_rows(payload["logic"]["id"])
            result = compare_records(sap_rows, logic_rows, payload["sap"]["mapping"], payload["logic"]["mapping"])
            self.server.last_result = result
            self.send_json(result)
        except Exception as exc:
            self.send_json({"error": f"对比失败：{exc}"}, HTTPStatus.BAD_REQUEST)

    def handle_export(self, table_name):
        try:
            result = getattr(self.server, "last_result", None)
            if not result:
                self.send_json({"error": "还没有可导出的对比结果"}, HTTPStatus.BAD_REQUEST)
                return
            rows = result.get(table_name, [])
            filename = f"{table_name}.csv"
            self.send_text(make_csv(rows), "text/csv; charset=utf-8", filename=filename)
        except Exception as exc:
            self.send_json({"error": f"导出失败：{exc}"}, HTTPStatus.BAD_REQUEST)

    def handle_save_export(self, table_name):
        try:
            result = getattr(self.server, "last_result", None)
            if not result:
                self.send_json({"error": "还没有可保存的对比结果"}, HTTPStatus.BAD_REQUEST)
                return
            rows = result.get(table_name, [])
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = safe_filename(f"{table_name}_{timestamp}.csv")
            path = EXPORT_DIR / filename
            path.write_text(make_csv(rows), encoding="utf-8-sig")
            self.send_json({"ok": True, "path": str(path), "rows": len(rows)})
        except Exception as exc:
            self.send_json({"error": f"保存失败：{exc}"}, HTTPStatus.BAD_REQUEST)


def run():
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "8765"))
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"Inventory compare tool running at http://{host}:{port}")
    server.serve_forever()


if __name__ == "__main__":
    run()
